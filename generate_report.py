import os
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_report():
    doc = Document()
    
    # Title Page
    for i in range(5): doc.add_paragraph() # Spacer
    title = doc.add_heading('Project Report: Predictive Modeling of Housing Price Dynamics', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    subtitle = doc.add_paragraph('Analyzing the Zillow Home Value Index (ZHVI) for US Metropolitan Areas')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    author = doc.add_paragraph('\n\nCourse: ADY-2026 | FPT University')
    author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_page_break()
    
    # 1. Introduction & overview of data science context
    doc.add_heading('1. Introduction', level=1)
    doc.add_paragraph(
        "The global real estate market is a critical pillar of economic infrastructure. For individual homeowners, property "
        "often represents their most significant financial asset. For institutions, real estate serves as a hedge against inflation "
        "and a cornerstone for long-term investment portfolios. However, the inherent complexity of housing markets—driven by "
        "interest rates, regional supply-demand imbalances, and psychological sentiment—makes price forecasting a notoriously "
        "difficult task."
    )
    doc.add_paragraph(
        "Traditional house price models often predict absolute price levels, which frequently leads to non-stationarity issues "
        "and high autocorrelation. This project addresses these challenges by shifting the modeling objective from absolute prices "
        "to stationary 'log-returns' (delta-log-price). By predicting the rate of growth rather than the price itself, our "
        "methodology ensures that the models remain robust across cities with vastly different price scales (e.g., San Francisco vs. Detroit)."
    )
    doc.add_paragraph(
        "Our primary objective is to evaluate whether linear models like Ridge Regression can effectively capture market momentum "
        "signals when benchmarked against non-linear architectures like LightGBM and simple zero-growth persistence models. "
        "The project utilizes the Zillow Home Value Index (ZHVI), which provides a stable, seasonally adjusted view of the US housing market."
    )
    
    # 2. Methodology (workflow, tools, libraries)
    doc.add_heading('2. Methodology', level=1)
    doc.add_paragraph(
        "The development pipeline is designed for scalability and reproducibility, leveraging modern data engineering practices "
        "to ensure a leakage-proof evaluation."
    )
    
    doc.add_heading('2.1 Data Pipeline and Workflow', level=2)
    doc.add_paragraph(
        "The workflow consists of five distinct phases:\n"
        "1. Ingestion: Raw wide-format CSV files containing monthly ZHVI data are melted into a long-format structure [RegionID, Date, Price]. "
        "This transformation enables per-region time-series analysis.\n"
        "2. Preprocessing: We apply time-aware linear interpolation to fill gaps in regions with sparse reporting. The target variable "
        "is computed as: log_return_t = log(Price_t / Price_{t-1}). This ensures the model predicts changes in market speed.\n"
        "3. Feature Engineering: Momentum features are captured using multi-interval lags (1, 2, 3, 6, and 12 months). This captures "
        "both short-term trend following and annual seasonality.\n"
        "4. City Target Encoding: To incorporate regional market identity, we encode each RegionID as its average historical log-return "
        "computed on the training set. This serves as a 'structural growth baseline' for each metro.\n"
        "5. Normalization: Features are standardized using StandardScaler fitted exclusively on the training corpus."
    )
    
    doc.add_heading('2.2 Model Selection and Hyperparameters', level=2)
    doc.add_paragraph(
        "We compare three distinct modeling philosophies:\n"
        "- Zero Growth Baseline: A simple 'persistence' assumption that tomorrow's price change is zero. This provides a sanity "
        "check for model performance.\n"
        "- Ridge Regression: A linear model with L2 regularization. We use TimeSeriesSplit for cross-validation to tune the alpha "
        "parameter, ensuring the model generalizes well to future time periods.\n"
        "- LightGBM Regressor: A gradient-boosted tree model that can capture non-linear interactions between lags. We utilize "
        "a fixed 15% validation window for early stopping to prevent overfitting."
    )
    
    doc.add_heading('2.3 Tools and Libraries', level=2)
    doc.add_paragraph(
        "The implementation is built on the Python 3.12 ecosystem, utilizing:\n"
        "- Pandas & NumPy for data manipulation.\n"
        "- Scikit-Learn for Ridge modeling, cross-validation, and scaling.\n"
        "- LightGBM for non-linear regression analysis.\n"
        "- Matplotlib & Seaborn for visualizing error residuals and feature importance."
    )
    
    # 3. Results: tables + visualizations + regression
    doc.add_heading('3. Results', level=1)
    doc.add_paragraph(
        "Overall performance was outstanding, with all trained models significantly outperforming the zero-growth baseline. "
        "The results are summarized in the table below, comparing predictions in both the logarithmic target space and the "
        "original USD currency space."
    )
    
    # Results Table
    table = doc.add_table(rows=1, cols=4)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Model Architecture'
    hdr_cells[1].text = 'log_return MAE'
    hdr_cells[2].text = 'Price MAE ($)'
    hdr_cells[3].text = 'Price R²'
    
    data = [
        ('Smart Baseline (Zero Growth)', '0.006913', '$1,795.06', '0.999668'),
        ('Ridge Regression (L2)', '0.002009', '$495.26', '0.999976'),
        ('LightGBM (Gradient Boosted)', '0.002166', '$543.87', '0.999967')
    ]
    
    for model, log_mae, price_mae, price_r2 in data:
        row_cells = table.add_row().cells
        row_cells[0].text = model
        row_cells[1].text = log_mae
        row_cells[2].text = price_mae
        row_cells[3].text = price_r2

    doc.add_paragraph("") # Space
    doc.add_paragraph(
        "Analysis of Error Metrics:\n"
        "- MAE ($): Ridge Regression achieves an incredible $495 average error across all regions. This represents a >70% gain in "
        "precision compared to the zero-growth assumption ($1,795 MAE).\n"
        "- R² Correlation: The R² values exceeding 0.999 reflect the model's ability to track the underlying structural trend of "
        "house prices. While R² on price levels is typically high due to momentum, the improvement in log-return MAE proves that "
        "the model is effectively predicting fluctuations in market velocity."
    )
    
    doc.add_heading('3.1 Feature Importance Analysis', level=2)
    doc.add_paragraph(
        "Feature importance logs reveal that the 1-month lag (`lag_1`) is the most dominant predictor, followed by `lag_2`. "
        "This confirms that the US housing market is highly momentum-driven. The `city_enc` feature also plays a significant "
        "role, accounting for the baseline growth rate differences between high-growth markets like Phoenix and stable markets "
        "like Cleveland."
    )
    
    # 4. Discussion: what the results mean in the real-world context
    doc.add_heading('4. Discussion', level=1)
    doc.add_heading('4.1 Linear vs. Non-Linear Dominance', level=2)
    doc.add_paragraph(
        "A surprising finding was that Ridge Regression slightly outperformed LightGBM. This suggests that the relationship "
        "between previous monthly returns and short-term future returns is predominantly linear. Real estate market friction—high "
        "transaction costs and long closing times—acts as a smoothing filter, creating stable trends that linear models are "
        "highly adept at capturing without the risk of overfitting inherent in deep trees."
    )
    doc.add_heading('4.2 Real-World Application and Insights', level=2)
    doc.add_paragraph(
        "For stakeholders (homebuyers and investors), this model provides a highly reliable 'momentum signal'. An average error "
        "of ~$500 on assets worth several hundred thousand dollars means the model can identify region-specific market cooling "
        "or acceleration months before they become obvious in raw price charts."
    )
    doc.add_heading('4.3 Limitations and Edge Cases', level=2)
    doc.add_paragraph(
        "The primary limitation is the model's reliance on backward-looking momentum. While excellent for forecasting continuous "
        "trends, the model cannot predict external economic shocks (e.g., pandemic lockdowns, sudden mortgage rate spikes) "
        "that have not yet appeared in the transactional data. Future iterations could incorporate exogenous variables such as "
        "the 30-year fixed mortgage rate or regional unemployment data."
    )
    
    doc.add_page_break()
    
    # Presentation Outline section (Bonus)
    doc.add_heading('Presentation Outline (10 Slides)', level=1)
    slides = [
        ("Slide 1: Background & Problem Context", "Overview of the US housing market and the challenge of non-stationary price prediction."),
        ("Slide 2: Research Question & Objectives", "Q: Can log-return momentum reliably forecast house prices? Obj: Build a high-precision Zillow-based pipeline."),
        ("Slide 3: Dataset Overview", "ZHVI for US MSAs. Wide-to-long transformation. Time-aware interpolation for data quality."),
        ("Slide 4: Data Preparation & Workflow", "Melt -> Interpolate -> Stationarity Transform (Log-Return) -> Time-Split."),
        ("Slide 5: Methodology", "Feature engineering (Lags t-1...t-12) and target-encoding for city momentum identity."),
        ("Slide 6: EDA & Statistics", "Identifying high autocorrelation (0.912 at lag-1). Log-return distribution analysis."),
        ("Slide 7: Modelling Strategy", "Linear models (Ridge) vs Non-linear (LightGBM). GridSearchCV for regularisation tuning."),
        ("Slide 8: Key Results & Visualizations", "Performance comparison: Ridge dominates with $495 Average Absolute Error."),
        ("Slide 9: Discussion & Insights", "Market inertia is stable. Linear models outperform trees due to reduced overfitting on smoothed data."),
        ("Slide 10: Conclusion & AI Usage", "Successful delivery of a production-ready model. Reflection on AI assistance for audit and validation.")
    ]
    
    for s_title, s_content in slides:
        p = doc.add_paragraph(style='List Bullet')
        run = p.add_run(s_title + ": ")
        run.bold = True
        p.add_run(s_content)

    doc.save('Mini_Report.docx')
    print("Report created successfully!")

if __name__ == "__main__":
    create_report()
