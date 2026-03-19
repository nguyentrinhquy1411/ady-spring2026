import os
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_presentation():
    doc = Document()
    
    # Title Page
    for i in range(5): doc.add_paragraph() # Spacer
    title = doc.add_heading('Presentation Outline: Predictive Modeling of Housing Price Dynamics', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    subtitle = doc.add_paragraph('Analyzing the Zillow Home Value Index (ZHVI) for US Metropolitan Areas')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    author = doc.add_paragraph('\n\nCourse: ADY-2026 | FPT University')
    author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_page_break()
    
    slides = [
        ("Slide 1: Background & Problem Context", [
            "Overview of the US housing market and its economic importance.",
            "The challenge of predicting non-stationary price levels.",
            "Why shifting to log-returns (market momentum) is more effective."
        ]),
        ("Slide 2: Research Question & Objectives", [
            "Can short-term log-return momentum reliably forecast house prices?",
            "Objective: Build a high-precision, leakage-proof time-series pipeline.",
            "Goal: Achieve < $500 Average Absolute Error across MSAs."
        ]),
        ("Slide 3: Dataset Overview", [
            "Source: Zillow Home Value Index (ZHVI) for US MSAs.",
            "Data Structure: Wide-format time series (Region per row, Date per column).",
            "Initial step: Wide-to-long transformation for chronological analysis."
        ]),
        ("Slide 4: Data Preparation & Workflow", [
            "Time-aware linear interpolation used for data quality and gap filling.",
            "Target Transformation: log(Price_t / Price_{t-1}) applied to achieve stationarity.",
            "Chronological splitting (80% Train, 20% Test) to prevent future data leakage."
        ]),
        ("Slide 5: Methodology", [
            "Feature Engineering: Extracting Lags at t-1, t-2, t-3, t-6, and t-12 months.",
            "City Target-Encoding: Capturing baseline momentum identity per region (computed only on train set).",
            "Scaling: StandardScaler fitted exclusively on training data for features."
        ]),
        ("Slide 6: EDA & Statistics", [
            "High autocorrelation (0.912) detected at lag-1 across markets.",
            "Log-returns distribution analysis showed near-zero centering (~0.003-0.005/month).",
            "Market inertia is the most significant observable predictor."
        ]),
        ("Slide 7: Modelling Strategy", [
            "Evaluated against a Smart Baseline (Zero-Growth assumption).",
            "Linear Model: Ridge Regression (L2 regularization) with TimeSeriesSplit.",
            "Non-Linear Model: LightGBM (Gradient Boosted Trees) with early stopping on validation set."
        ]),
        ("Slide 8: Key Results & Visualizations", [
            "Ridge Regression dominates with $495 Average Absolute Error in price space.",
            "This represents a >70% error reduction vs Zero-Growth Baseline ($1,795 MAE).",
            "R2 values exceeding 0.999 demonstrate strong tracking of the underlying structural trend."
        ]),
        ("Slide 9: Discussion & Insights", [
            "Market inertia is stable. Linear models outperform complex trees by avoiding overfitting on smoothed data.",
            "The pipeline provides a reliable 'momentum signal' for investors before trends are obvious.",
            "Limitations include reliance on backward-looking momentum rather than exogenous economic shocks."
        ]),
        ("Slide 10: Conclusion & AI Usage", [
            "Successful delivery of a production-ready model for ZHVI forecasting.",
            "AI Assistance: System architecture design, leakage prevention auditing, and metrics formulation.",
            "Rigorous manual and statistical validation employed to ensure pipeline integrity."
        ])
    ]
    
    for s_title, s_content_list in slides:
        doc.add_heading(s_title, level=1)
        for point in s_content_list:
            p = doc.add_paragraph(style='List Bullet')
            p.add_run(point)
        doc.add_page_break()

    # Save to the root project directory so user can easily find it
    # __file__ is scripts/generate_presentation.py so dirname dirname is root
    out_path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'Presentation.docx')
    doc.save(out_path)
    print(f"Presentation created successfully at: {out_path}")

if __name__ == "__main__":
    create_presentation()
